"""Строгий разбор Excel-файла FAQ до любых изменений базы."""

from __future__ import annotations

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from app.db.models import FaqItem, FaqType
from app.services.faq_excel import (
    TEMPLATE_VERSION,
    FaqWorkbookError,
    build_faq_workbook,
    parse_faq_workbook,
)

TYPE_HEADERS = ["ID типа", "Название", "Порядок", "Действие"]
QUESTION_HEADERS = [
    "ID вопроса",
    "ID типа",
    "Вопрос",
    "Ответ",
    "Порядок",
    "Действие",
]


def workbook_bytes(
    *,
    types: list[list[object]] | None = None,
    questions: list[list[object]] | None = None,
    omit_sheet: str | None = None,
    type_headers: list[str] | None = None,
    template_version: int = 2,
) -> bytes:
    workbook = Workbook()
    instruction = workbook.active
    instruction.title = "Инструкция"
    instruction.append(["FAQ_TEMPLATE_VERSION", template_version])
    instruction.append(["FAQ_NEXT_TYPE_ID", 1])

    type_sheet = workbook.create_sheet("Типы")
    type_sheet.append(type_headers or TYPE_HEADERS)
    for row in types or []:
        type_sheet.append(row)

    question_sheet = workbook.create_sheet("Вопросы")
    question_sheet.append(QUESTION_HEADERS)
    for row in questions or []:
        question_sheet.append(row)

    if omit_sheet:
        del workbook[omit_sheet]
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def without_instruction_dimension(payload: bytes) -> bytes:
    """Имитировать XLSX редактора, который не записывает dimension листа."""
    source = BytesIO(payload)
    output = BytesIO()
    with ZipFile(source) as archive, ZipFile(output, "w", ZIP_DEFLATED) as rewritten:
        for entry in archive.infolist():
            contents = archive.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                import re

                contents = re.sub(rb"<dimension[^>]*/>", b"", contents, count=1)
            rewritten.writestr(entry, contents)
    return output.getvalue()


def test_parser_accepts_copy_pasted_actions_and_integer_links() -> None:
    """Шаблоны версии 2 с заголовками в первой строке остаются рабочими."""
    payload = workbook_bytes(
        types=[[8, "Доставка", 10, "сохранить"]],
        questions=[[None, 8, "Когда?", "Завтра", 20, "СОХРАНИТЬ"]],
    )
    parsed = parse_faq_workbook(payload)
    assert parsed.types[0].id == 8
    assert parsed.questions[0].type_id == 8


def test_parser_accepts_workbook_without_instruction_sheet_dimension() -> None:
    payload = without_instruction_dimension(
        workbook_bytes(types=[[8, "Доставка", 10, "СОХРАНИТЬ"]])
    )

    parsed = parse_faq_workbook(payload)

    assert parsed.template_next_type_id == 1
    assert parsed.types[0].title == "Доставка"


@pytest.mark.parametrize("action", ["", "ДОБАВИТЬ", "DELETE", "СОХРАНИТЬ "])
def test_parser_rejects_unknown_action_with_sheet_and_row(action: str) -> None:
    with pytest.raises(FaqWorkbookError, match=r"Типы.*строка 2.*Действие"):
        parse_faq_workbook(workbook_bytes(types=[[8, "Доставка", 10, action]]))


def test_parser_requires_all_sheets_and_exact_headers() -> None:
    with pytest.raises(FaqWorkbookError, match="Вопросы"):
        parse_faq_workbook(workbook_bytes(omit_sheet="Вопросы"))
    with pytest.raises(FaqWorkbookError, match="заголов"):
        parse_faq_workbook(workbook_bytes(type_headers=["id", *TYPE_HEADERS[1:]]))


@pytest.mark.parametrize("bad_id", [1.5, "8", 0, -1, True])
def test_parser_rejects_non_positive_integer_ids(bad_id: object) -> None:
    with pytest.raises(FaqWorkbookError, match="ID типа"):
        parse_faq_workbook(workbook_bytes(types=[[bad_id, "Тип", 0, "СОХРАНИТЬ"]]))


def test_parser_rejects_duplicate_ids() -> None:
    with pytest.raises(FaqWorkbookError, match="повтор"):
        parse_faq_workbook(
            workbook_bytes(
                types=[
                    [8, "Один", 0, "СОХРАНИТЬ"],
                    [8, "Два", 1, "СОХРАНИТЬ"],
                ]
            )
        )


def test_parser_rejects_empty_or_too_long_text() -> None:
    with pytest.raises(FaqWorkbookError, match="Название"):
        parse_faq_workbook(workbook_bytes(types=[[8, "", 0, "СОХРАНИТЬ"]]))
    with pytest.raises(FaqWorkbookError, match="Ответ"):
        parse_faq_workbook(
            workbook_bytes(
                types=[[8, "Тип", 0, "СОХРАНИТЬ"]],
                questions=[[None, 8, "Вопрос", "x" * 4001, 0, "СОХРАНИТЬ"]],
            )
        )


@pytest.mark.parametrize("formula", ["=1+1", "='[other.xlsx]Sheet1'!A1"])
def test_parser_rejects_formulas_and_external_links(formula: str) -> None:
    with pytest.raises(FaqWorkbookError, match="формул"):
        parse_faq_workbook(workbook_bytes(types=[[8, formula, 0, "СОХРАНИТЬ"]]))


def test_parser_handles_existing_orphan_but_rejects_new_orphan() -> None:
    parsed = parse_faq_workbook(
        workbook_bytes(questions=[[7, None, "Старый?", "Да", 0, "СОХРАНИТЬ"]])
    )
    assert parsed.questions[0].type_id is None

    with pytest.raises(FaqWorkbookError, match="нового вопроса"):
        parse_faq_workbook(workbook_bytes(questions=[[None, None, "Новый?", "Да", 0, "СОХРАНИТЬ"]]))


def test_parser_rejects_unknown_or_deleted_type_reference() -> None:
    with pytest.raises(FaqWorkbookError, match="неизвестный тип"):
        parse_faq_workbook(workbook_bytes(questions=[[None, 8, "Вопрос", "Ответ", 0, "СОХРАНИТЬ"]]))
    with pytest.raises(FaqWorkbookError, match="удаляемый тип"):
        parse_faq_workbook(
            workbook_bytes(
                types=[[8, "Тип", 0, "УДАЛИТЬ"]],
                questions=[[None, 8, "Вопрос", "Ответ", 0, "СОХРАНИТЬ"]],
            )
        )


def test_parser_skips_id_only_prepared_type_rows() -> None:
    parsed = parse_faq_workbook(workbook_bytes(types=[[8, None, None, None]]))
    assert parsed.types == ()
    assert parsed.prepared_type_ids == (8,)


def test_parser_enforces_row_and_file_limits() -> None:
    type_rows = [[value, f"Тип {value}", 0, "СОХРАНИТЬ"] for value in range(1, 502)]
    with pytest.raises(FaqWorkbookError, match="500"):
        parse_faq_workbook(workbook_bytes(types=type_rows))
    with pytest.raises(FaqWorkbookError, match="5 МБ"):
        parse_faq_workbook(b"x" * (5 * 1024 * 1024 + 1))


def test_export_contains_instruction_current_data_and_prepared_type_ids() -> None:
    payload = build_faq_workbook(
        types=[FaqType(id=7, title="Доставка", sort_order=10, is_active=True)],
        questions=[
            FaqItem(
                id=12,
                type_id=7,
                question="Когда?",
                answer="Завтра",
                sort_order=5,
                is_active=True,
            ),
            FaqItem(
                id=13,
                type_id=None,
                question="Сирота?",
                answer="Да",
                sort_order=6,
                is_active=True,
            ),
        ],
        next_type_id=8,
        new_type_slots=3,
    )
    book = load_workbook(BytesIO(payload), data_only=True)
    assert book.sheetnames == ["Инструкция", "Типы", "Вопросы"]
    assert book["Инструкция"]["A1"].value == "Как заполнить шаблон FAQ"
    assert "ID типа" in " ".join(
        str(book["Инструкция"].cell(row=row, column=1).value or "") for row in range(1, 12)
    )
    assert [book["Типы"][f"A{row}"].value for row in range(3, 6)] == [8, 9, 10]
    assert [cell.value for cell in book["Типы"][1]] == TYPE_HEADERS
    question_sheet = book["Вопросы"]
    assert "для нового вопроса" in str(question_sheet["B1"].value).lower()
    assert "оставьте" in str(question_sheet["B1"].value).lower()
    assert "id вопроса" in str(question_sheet["B1"].value).lower()
    assert {str(value) for value in question_sheet.merged_cells.ranges} == {"B1:F1"}
    assert book["Инструкция"]["B2"].value == TEMPLATE_VERSION
    assert [cell.value for cell in question_sheet[2]] == QUESTION_HEADERS
    assert [book["Типы"].cell(2, col).value for col in range(1, 5)] == [
        7,
        "Доставка",
        10,
        "СОХРАНИТЬ",
    ]
    assert question_sheet["B4"].value is None
    assert book["Типы"].freeze_panes == "A2"
    assert question_sheet.freeze_panes == "A3"
    assert question_sheet.column_dimensions["A"].hidden is True
    assert question_sheet.auto_filter.ref is not None
    validations = list(book["Типы"].data_validations.dataValidation)
    assert validations and validations[0].formula1 == '"СОХРАНИТЬ,УДАЛИТЬ"'


def test_exported_workbook_round_trips_after_adding_type_and_question() -> None:
    payload = build_faq_workbook(
        types=[FaqType(id=7, title="Доставка", sort_order=10, is_active=True)],
        questions=[],
        next_type_id=8,
        new_type_slots=2,
    )
    book = load_workbook(BytesIO(payload))
    type_sheet = book["Типы"]
    type_sheet["B3"] = "Оплата"
    type_sheet["C3"] = 20
    type_sheet["D3"] = "СОХРАНИТЬ"
    book["Вопросы"].append([None, 8, "Как?", "Так", 0, "СОХРАНИТЬ"])
    output = BytesIO()
    book.save(output)

    parsed = parse_faq_workbook(output.getvalue())
    assert [row.id for row in parsed.types] == [7, 8]
    assert parsed.questions[0].type_id == 8
    assert parsed.prepared_type_ids == (9,)
