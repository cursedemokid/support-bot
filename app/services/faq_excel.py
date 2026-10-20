"""Безопасный формат обмена типами и вопросами FAQ через XLSX."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any, cast
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select, text, update
from sqlalchemy.engine import CursorResult
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import FaqItem, FaqType

SAVE = "СОХРАНИТЬ"
DELETE = "УДАЛИТЬ"
TEMPLATE_VERSION = 3
SUPPORTED_TEMPLATE_VERSIONS = frozenset({2, TEMPLATE_VERSION})
MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_TYPE_ROWS = 500
MAX_QUESTION_ROWS = 2_000

TYPE_SHEET = "Типы"
QUESTION_SHEET = "Вопросы"
INSTRUCTION_SHEET = "Инструкция"
TYPE_HEADERS = ("ID типа", "Название", "Порядок", "Действие")
QUESTION_HEADERS = (
    "ID вопроса",
    "ID типа",
    "Вопрос",
    "Ответ",
    "Порядок",
    "Действие",
)
QUESTION_NOTICE = (
    "Колонка «ID вопроса» скрыта: для нового вопроса оставьте ID пустым — бот назначит "
    "его автоматически. ID нужен только для существующих вопросов."
)


@dataclass(frozen=True, slots=True)
class FaqWorkbookError(ValueError):
    sheet: str
    row: int | None
    column: str | None
    detail: str

    def __str__(self) -> str:
        location = f"Лист «{self.sheet}»"
        if self.row is not None:
            location += f", строка {self.row}"
        if self.column is not None:
            location += f", столбец «{self.column}»"
        return f"{location}: {self.detail}"


@dataclass(frozen=True, slots=True)
class FaqTypeRow:
    row: int
    id: int
    title: str
    sort_order: int
    action: str


@dataclass(frozen=True, slots=True)
class FaqQuestionRow:
    row: int
    id: int | None
    type_id: int | None
    question: str
    answer: str
    sort_order: int
    action: str


@dataclass(frozen=True, slots=True)
class FaqWorkbookData:
    types: tuple[FaqTypeRow, ...]
    questions: tuple[FaqQuestionRow, ...]
    prepared_type_ids: tuple[int, ...]
    template_next_type_id: int | None = None


@dataclass(frozen=True, slots=True)
class FaqImportSummary:
    created_types: int = 0
    updated_types: int = 0
    deleted_types: int = 0
    created_questions: int = 0
    updated_questions: int = 0
    deleted_questions: int = 0
    orphaned_questions: int = 0


def format_import_summary(summary: FaqImportSummary) -> str:
    return (
        "FAQ обновлён.\n"
        f"Типы: добавлено {summary.created_types}, "
        f"изменено {summary.updated_types}, удалено {summary.deleted_types}.\n"
        f"Вопросы: добавлено {summary.created_questions}, "
        f"изменено {summary.updated_questions}, "
        f"удалено {summary.deleted_questions}.\n"
        f"Без типа: {summary.orphaned_questions}."
    )


def _error(
    sheet: str, detail: str, *, row: int | None = None, column: str | None = None
) -> FaqWorkbookError:
    return FaqWorkbookError(sheet=sheet, row=row, column=column, detail=detail)


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _ensure_plain(cell: Any, *, sheet: str, row: int, column: str) -> None:
    value = cell.value
    if cell.data_type in {"f", "e"} or (isinstance(value, str) and value.startswith("=")):
        raise _error(sheet, "формулы и внешние ссылки запрещены", row=row, column=column)


def _positive_int(value: object, *, sheet: str, row: int, column: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
        raise _error(sheet, "ожидается положительное целое число", row=row, column=column)
    return value


def _integer(value: object, *, sheet: str, row: int, column: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int):
        raise _error(sheet, "ожидается целое число", row=row, column=column)
    return value


def _text(
    value: object,
    *,
    sheet: str,
    row: int,
    column: str,
    max_length: int,
) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > max_length:
        raise _error(
            sheet,
            f"ожидается текст от 1 до {max_length} символов",
            row=row,
            column=column,
        )
    return value


def _action(value: object, *, sheet: str, row: int) -> str:
    if not isinstance(value, str) or value != value.strip():
        raise _error(sheet, f"допустимы только {SAVE} и {DELETE}", row=row, column="Действие")
    normalized = value.upper()
    if normalized not in {SAVE, DELETE}:
        raise _error(sheet, f"допустимы только {SAVE} и {DELETE}", row=row, column="Действие")
    return normalized


def _check_headers(sheet: Any, expected: tuple[str, ...], *, row: int = 1) -> None:
    actual = tuple(
        cell.value
        for cell in next(sheet.iter_rows(min_row=row, max_row=row, max_col=len(expected)))
    )
    if actual != expected:
        raise _error(sheet.title, f"заголовки должны быть: {', '.join(expected)}", row=row)


def parse_faq_workbook(payload: bytes) -> FaqWorkbookData:
    if len(payload) > MAX_FILE_SIZE:
        raise _error("Файл", "размер не должен превышать 5 МБ")
    try:
        workbook = load_workbook(
            BytesIO(payload), read_only=True, data_only=False, keep_links=False
        )
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise _error("Файл", "не удалось прочитать XLSX") from exc

    try:
        for required in (INSTRUCTION_SHEET, TYPE_SHEET, QUESTION_SHEET):
            if required not in workbook.sheetnames:
                raise _error(required, "обязательный лист отсутствует")
        instruction = workbook[INSTRUCTION_SHEET]
        # Некоторые редакторы сохраняют корректный лист без XML-тега
        # ``dimension``. В read-only режиме openpyxl тогда возвращает
        # ``max_row=None``, поэтому читаем фиксированный безопасный диапазон,
        # не полагаясь на метаданные размеров листа.
        instruction_metadata = {
            key: value
            for key, value in instruction.iter_rows(
                min_row=1,
                max_row=10,
                min_col=1,
                max_col=2,
                values_only=True,
            )
            if isinstance(key, str)
        }
        version = instruction_metadata.get("FAQ_TEMPLATE_VERSION")
        if version not in SUPPORTED_TEMPLATE_VERSIONS:
            raise _error(INSTRUCTION_SHEET, "неподдерживаемая версия шаблона", row=1)
        template_next_type_id = instruction_metadata.get("FAQ_NEXT_TYPE_ID")
        if template_next_type_id is None:
            raise _error(INSTRUCTION_SHEET, "не указан следующий ID типа", row=3)
        template_next_type_id = _positive_int(
            template_next_type_id,
            sheet=INSTRUCTION_SHEET,
            row=3,
            column="FAQ_NEXT_TYPE_ID",
        )

        type_sheet = workbook[TYPE_SHEET]
        question_sheet = workbook[QUESTION_SHEET]
        _check_headers(type_sheet, TYPE_HEADERS)
        question_header_row = 2 if version == TEMPLATE_VERSION else 1
        _check_headers(question_sheet, QUESTION_HEADERS, row=question_header_row)
        types, prepared_ids = _parse_types(type_sheet)
        questions = _parse_questions(question_sheet, header_row=question_header_row)
        _validate_references(types, questions)
        return FaqWorkbookData(
            types=tuple(types),
            questions=tuple(questions),
            prepared_type_ids=tuple(prepared_ids),
            template_next_type_id=template_next_type_id,
        )
    finally:
        workbook.close()


def _parse_types(sheet: Any) -> tuple[list[FaqTypeRow], list[int]]:
    parsed: list[FaqTypeRow] = []
    prepared: list[int] = []
    seen_ids: set[int] = set()
    meaningful_rows = 0
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=4), start=2):
        for cell, column in zip(cells, TYPE_HEADERS, strict=True):
            _ensure_plain(cell, sheet=TYPE_SHEET, row=row_number, column=column)
        values = [cell.value for cell in cells]
        if all(_is_blank(value) for value in values):
            continue
        meaningful_rows += 1
        if meaningful_rows > MAX_TYPE_ROWS:
            raise _error(TYPE_SHEET, f"не более {MAX_TYPE_ROWS} строк", row=row_number)
        type_id = _positive_int(values[0], sheet=TYPE_SHEET, row=row_number, column="ID типа")
        if type_id in seen_ids:
            raise _error(TYPE_SHEET, f"повтор ID {type_id}", row=row_number, column="ID типа")
        seen_ids.add(type_id)
        if all(_is_blank(value) for value in values[1:]):
            prepared.append(type_id)
            continue
        parsed.append(
            FaqTypeRow(
                row=row_number,
                id=type_id,
                title=_text(
                    values[1], sheet=TYPE_SHEET, row=row_number, column="Название", max_length=100
                ),
                sort_order=_integer(values[2], sheet=TYPE_SHEET, row=row_number, column="Порядок"),
                action=_action(values[3], sheet=TYPE_SHEET, row=row_number),
            )
        )
    return parsed, prepared


def _parse_questions(sheet: Any, *, header_row: int) -> list[FaqQuestionRow]:
    parsed: list[FaqQuestionRow] = []
    seen_ids: set[int] = set()
    meaningful_rows = 0
    first_data_row = header_row + 1
    for row_number, cells in enumerate(
        sheet.iter_rows(min_row=first_data_row, max_col=6), start=first_data_row
    ):
        for cell, column in zip(cells, QUESTION_HEADERS, strict=True):
            _ensure_plain(cell, sheet=QUESTION_SHEET, row=row_number, column=column)
        values = [cell.value for cell in cells]
        if all(_is_blank(value) for value in values):
            continue
        meaningful_rows += 1
        if meaningful_rows > MAX_QUESTION_ROWS:
            raise _error(QUESTION_SHEET, f"не более {MAX_QUESTION_ROWS} строк", row=row_number)

        question_id = None
        if not _is_blank(values[0]):
            question_id = _positive_int(
                values[0], sheet=QUESTION_SHEET, row=row_number, column="ID вопроса"
            )
            if question_id in seen_ids:
                raise _error(
                    QUESTION_SHEET,
                    f"повтор ID {question_id}",
                    row=row_number,
                    column="ID вопроса",
                )
            seen_ids.add(question_id)
        type_id = None
        if not _is_blank(values[1]):
            type_id = _positive_int(
                values[1], sheet=QUESTION_SHEET, row=row_number, column="ID типа"
            )
        action = _action(values[5], sheet=QUESTION_SHEET, row=row_number)
        if question_id is None and action == DELETE:
            raise _error(
                QUESTION_SHEET, "для удаления нужен ID вопроса", row=row_number, column="ID вопроса"
            )
        if question_id is None and type_id is None:
            raise _error(
                QUESTION_SHEET, "для нового вопроса нужен ID типа", row=row_number, column="ID типа"
            )
        parsed.append(
            FaqQuestionRow(
                row=row_number,
                id=question_id,
                type_id=type_id,
                question=_text(
                    values[2], sheet=QUESTION_SHEET, row=row_number, column="Вопрос", max_length=500
                ),
                answer=_text(
                    values[3], sheet=QUESTION_SHEET, row=row_number, column="Ответ", max_length=4000
                ),
                sort_order=_integer(
                    values[4], sheet=QUESTION_SHEET, row=row_number, column="Порядок"
                ),
                action=action,
            )
        )
    return parsed


def _validate_references(types: list[FaqTypeRow], questions: list[FaqQuestionRow]) -> None:
    type_actions = {row.id: row.action for row in types}
    for question in questions:
        if question.type_id is None:
            continue
        action = type_actions.get(question.type_id)
        if action is None:
            raise _error(
                QUESTION_SHEET,
                f"указан неизвестный тип {question.type_id}",
                row=question.row,
                column="ID типа",
            )
        if question.action == SAVE and action == DELETE:
            raise _error(
                QUESTION_SHEET,
                f"нельзя сохранить вопрос в удаляемый тип {question.type_id}",
                row=question.row,
                column="ID типа",
            )


def build_faq_workbook(
    *,
    types: list[FaqType],
    questions: list[FaqItem],
    next_type_id: int,
    new_type_slots: int = 50,
) -> bytes:
    """Сформировать редактируемый шаблон с актуальными FAQ и свободными ID."""
    if next_type_id <= 0 or new_type_slots < 0:
        raise ValueError("next_type_id must be positive and new_type_slots non-negative")

    workbook = Workbook()
    instruction = workbook.active
    if not isinstance(instruction, Worksheet):
        raise RuntimeError("workbook has no active worksheet")
    instruction.title = INSTRUCTION_SHEET
    type_sheet = workbook.create_sheet(TYPE_SHEET)
    question_sheet = workbook.create_sheet(QUESTION_SHEET)

    _write_instruction(instruction, next_type_id=next_type_id)
    _write_table_header(type_sheet, TYPE_HEADERS)
    for faq_type in types:
        type_sheet.append([faq_type.id, faq_type.title, faq_type.sort_order, SAVE])
    for prepared_id in range(next_type_id, next_type_id + new_type_slots):
        type_sheet.append([prepared_id, None, None, None])

    _write_question_notice(question_sheet)
    _write_table_header(question_sheet, QUESTION_HEADERS)
    for faq_question in questions:
        question_sheet.append(
            [
                faq_question.id,
                faq_question.type_id,
                faq_question.question,
                faq_question.answer,
                faq_question.sort_order,
                SAVE,
            ]
        )

    _format_data_sheet(
        type_sheet,
        widths=(14, 34, 14, 18),
        action_column="D",
        validation_end=MAX_TYPE_ROWS + 1,
    )
    _format_data_sheet(
        question_sheet,
        widths=(16, 14, 42, 64, 14, 18),
        action_column="F",
        validation_end=MAX_QUESTION_ROWS + 2,
        header_row=2,
    )
    question_sheet.column_dimensions["A"].hidden = True
    for row in question_sheet.iter_rows(min_row=3, min_col=3, max_col=4):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _write_instruction(sheet: Worksheet, *, next_type_id: int) -> None:
    sheet["A1"] = "Как заполнить шаблон FAQ"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet.merge_cells("A1:B1")
    sheet.append(["FAQ_TEMPLATE_VERSION", TEMPLATE_VERSION])
    sheet.append(["FAQ_NEXT_TYPE_ID", next_type_id])
    instructions = (
        "1. Не переименовывайте листы и столбцы.",
        f"2. {SAVE} — добавить или обновить строку.",
        f"3. {DELETE} — удалить строку. Отсутствие строки ничего не удаляет.",
        "4. Для нового типа используйте первый свободный ID типа на листе «Типы».",
        "5. Этот ID типа можно сразу вставить в новый вопрос.",
        "6. Для нового вопроса оставьте «ID вопроса» пустым. "
        "После загрузки скачайте свежий шаблон.",
        "7. Вопрос без типа не показывается клиентам.",
    )
    for text_value in instructions:
        sheet.append([text_value])
    sheet.column_dimensions["A"].width = 110
    sheet.column_dimensions["B"].width = 18
    for row in sheet.iter_rows():
        row[0].alignment = Alignment(vertical="top", wrap_text=True)


def _write_table_header(sheet: Worksheet, headers: tuple[str, ...]) -> None:
    sheet.append(list(headers))
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_question_notice(sheet: Worksheet) -> None:
    sheet.append([None, QUESTION_NOTICE])
    sheet.merge_cells("B1:F1")
    cell = sheet["B1"]
    cell.font = Font(bold=True, color="7F6000")
    cell.fill = PatternFill("solid", fgColor="FFF2CC")
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 42


def _format_data_sheet(
    sheet: Worksheet,
    *,
    widths: tuple[int, ...],
    action_column: str,
    validation_end: int,
    header_row: int = 1,
) -> None:
    first_data_row = header_row + 1
    sheet.freeze_panes = f"A{first_data_row}"
    last_column = get_column_letter(len(widths))
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{max(sheet.max_row, first_data_row)}"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    validation = DataValidation(type="list", formula1=f'"{SAVE},{DELETE}"', allow_blank=True)
    validation.error = f"Введите {SAVE} или {DELETE}"
    validation.errorTitle = "Неверное действие"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"{action_column}{first_data_row}:{action_column}{validation_end}")


async def apply_faq_workbook(session: AsyncSession, data: FaqWorkbookData) -> FaqImportSummary:
    """Проверить актуальность шаблона и применить его в текущей транзакции."""
    if not session.in_transaction():
        raise RuntimeError("apply_faq_workbook requires an active transaction")

    await session.execute(text("select pg_advisory_xact_lock(716201001)"))
    existing_types = {value.id: value for value in (await session.scalars(select(FaqType))).all()}
    existing_questions = {
        value.id: value for value in (await session.scalars(select(FaqItem))).all()
    }
    current_next_id = int(
        await session.scalar(select(func.coalesce(func.max(FaqType.id), 0) + 1)) or 1
    )

    if data.template_next_type_id is not None and data.template_next_type_id != current_next_id:
        first_prepared_row = next(
            (row for row in data.types if row.id >= data.template_next_type_id),
            None,
        )
        raise _error(
            TYPE_SHEET,
            "ID уже изменились; скачайте свежий шаблон",
            row=first_prepared_row.row if first_prepared_row else None,
            column="ID типа",
        )

    new_type_rows = [
        row for row in data.types if row.action == SAVE and row.id not in existing_types
    ]
    expected_ids = list(range(current_next_id, current_next_id + len(new_type_rows)))
    if [row.id for row in new_type_rows] != expected_ids:
        first = (
            new_type_rows[0]
            if new_type_rows
            else next((row for row in data.types if row.id not in existing_types), None)
        )
        raise _error(
            TYPE_SHEET,
            "ID уже изменились; скачайте свежий шаблон",
            row=first.row if first else None,
            column="ID типа",
        )

    deleting_type_ids = {row.id for row in data.types if row.action == DELETE}
    available_type_ids = {
        type_id
        for type_id, value in existing_types.items()
        if value.is_active and type_id not in deleting_type_ids
    } | {row.id for row in new_type_rows}

    for type_row in data.types:
        existing_type = existing_types.get(type_row.id)
        if type_row.action == DELETE and existing_type is None:
            raise _error(TYPE_SHEET, "тип не найден", row=type_row.row, column="ID типа")
        if existing_type is not None and not existing_type.is_active:
            raise _error(TYPE_SHEET, "тип уже удалён", row=type_row.row, column="ID типа")

    for question_row in data.questions:
        existing_question = (
            existing_questions.get(question_row.id) if question_row.id is not None else None
        )
        if question_row.id is not None and (
            existing_question is None or not existing_question.is_active
        ):
            detail = "вопрос не найден или уже удалён"
            if question_row.action == SAVE:
                detail += "; для нового вопроса очистите «ID вопроса» — бот назначит его сам"
            raise _error(
                QUESTION_SHEET,
                detail,
                row=question_row.row,
                column="ID вопроса",
            )
        if (
            question_row.action == SAVE
            and question_row.type_id is not None
            and question_row.type_id not in available_type_ids
        ):
            raise _error(
                QUESTION_SHEET,
                f"нельзя сохранить вопрос в неизвестный или удаляемый тип {question_row.type_id}",
                row=question_row.row,
                column="ID типа",
            )
        if question_row.action == SAVE and question_row.id is None and question_row.type_id is None:
            raise _error(
                QUESTION_SHEET,
                "для нового вопроса нужен ID типа",
                row=question_row.row,
                column="ID типа",
            )

    counters = {
        "created_types": 0,
        "updated_types": 0,
        "deleted_types": 0,
        "created_questions": 0,
        "updated_questions": 0,
        "deleted_questions": 0,
        "orphaned_questions": 0,
    }
    for type_row in data.types:
        if type_row.action != SAVE:
            continue
        existing_type = existing_types.get(type_row.id)
        if existing_type is None:
            existing_type = FaqType(
                id=type_row.id,
                title=type_row.title,
                sort_order=type_row.sort_order,
                is_active=True,
            )
            session.add(existing_type)
            existing_types[type_row.id] = existing_type
            counters["created_types"] += 1
        else:
            existing_type.title = type_row.title
            existing_type.sort_order = type_row.sort_order
            counters["updated_types"] += 1

    # SQLAlchemy не знает об ORM-связи между этими моделями, поэтому без
    # явного flush может отправить новый вопрос раньше его нового типа.
    # Flush остаётся частью той же транзакции и будет полностью откатан при
    # любой последующей ошибке импорта.
    if new_type_rows:
        await session.flush()

    for question_row in data.questions:
        if question_row.action != SAVE:
            continue
        existing_question = (
            existing_questions.get(question_row.id) if question_row.id is not None else None
        )
        if existing_question is None:
            session.add(
                FaqItem(
                    type_id=question_row.type_id,
                    question=question_row.question,
                    answer=question_row.answer,
                    sort_order=question_row.sort_order,
                    is_active=True,
                )
            )
            counters["created_questions"] += 1
        else:
            existing_question.type_id = question_row.type_id
            existing_question.question = question_row.question
            existing_question.answer = question_row.answer
            existing_question.sort_order = question_row.sort_order
            counters["updated_questions"] += 1

    for question_row in data.questions:
        if question_row.action != DELETE or question_row.id is None:
            continue
        existing_questions[question_row.id].is_active = False
        counters["deleted_questions"] += 1

    for type_row in data.types:
        if type_row.action != DELETE:
            continue
        existing_types[type_row.id].is_active = False
        result = cast(
            CursorResult[Any],
            await session.execute(
                update(FaqItem)
                .where(FaqItem.type_id == type_row.id, FaqItem.is_active)
                .values(type_id=None)
            ),
        )
        counters["deleted_types"] += 1
        counters["orphaned_questions"] += int(result.rowcount or 0)

    await session.flush()
    return FaqImportSummary(
        created_types=counters["created_types"],
        updated_types=counters["updated_types"],
        deleted_types=counters["deleted_types"],
        created_questions=counters["created_questions"],
        updated_questions=counters["updated_questions"],
        deleted_questions=counters["deleted_questions"],
        orphaned_questions=counters["orphaned_questions"],
    )
